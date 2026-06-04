"""
Central Water Commission (CWC) / India-WRIS data adapter.

CWC provides:
  - Water level (gauge) data
  - Discharge measurements
  - Sediment/suspended load
  - Some water quality parameters

Data can be obtained from:
  - https://indiawris.gov.in (India-WRIS portal)
  - CWC real-time data system
  - Downloaded Excel/CSV files

Typical Excel format:
  Station Name | Basin | Sub-Basin | State | Date | Water Level (m) | Discharge (cumec)
"""

import csv
import logging
from datetime import date
from pathlib import Path
from typing import Any, Dict, List, Optional

from .base import BaseAdapter
from ..config.schema import CanonicalRecord

logger = logging.getLogger(__name__)

# Column name patterns in CWC data
_CWC_COLUMN_MAP = {
    "water level":      ("water_level",   "m"),
    "gauge level":      ("water_level",   "m"),
    "stage":            ("water_level",   "m"),
    "discharge":        ("discharge",     "m³/s"),
    "flow":             ("discharge",     "m³/s"),
    "inflow":           ("discharge",     "m³/s"),
    "sediment":         ("sediment_load", "mg/L"),
    "suspended load":   ("sediment_load", "mg/L"),
    "temperature":      ("temperature",   "°C"),
    "conductivity":     ("conductivity",  "µS/cm"),
    "turbidity":        ("turbidity",     "NTU"),
    "ph":               ("ph",            ""),
    "dissolved oxygen": ("do",            "mg/L"),
}


class CWCAdapter(BaseAdapter):
    """Adapter for CWC / India-WRIS hydrological data."""

    source_id = "cwc"
    source_name = "Central Water Commission (CWC)"

    def __init__(
        self,
        data_path: Optional[Path] = None,
        output_dir: Optional[Path] = None,
    ):
        super().__init__(output_dir)
        self.data_path = data_path

    def _detect_format(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix in (".csv", ".txt"):
            return "csv"
        if suffix in (".xls", ".xlsx"):
            return "excel"
        raise ValueError(f"Unsupported file format: {suffix}")

    def fetch(self, **kwargs) -> Any:
        """Read CWC data from local CSV or Excel file."""
        path = kwargs.get("path") or self.data_path
        if not path:
            logger.warning("No CWC data path provided, skipping.")
            return []
        path = Path(path)
        if not path.exists():
            logger.warning(f"CWC file not found: {path}")
            return []

        fmt = self._detect_format(path)

        if fmt == "csv":
            rows = []
            with open(path, newline="", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                self._raw_columns = list(reader.fieldnames or [])
                for row in reader:
                    rows.append(dict(row))
            logger.info(f"Read {len(rows)} CWC rows from {path}")
            return rows

        elif fmt == "excel":
            try:
                import openpyxl
            except ImportError:
                logger.error("openpyxl required for Excel files: pip install openpyxl")
                return []

            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            rows = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                data = list(ws.iter_rows(values_only=True))
                if len(data) < 2:
                    continue
                headers = [str(h).strip() if h else "" for h in data[0]]
                self._raw_columns = headers
                for row_vals in data[1:]:
                    row_dict = {
                        h: (str(v).strip() if v is not None else "")
                        for h, v in zip(headers, row_vals)
                    }
                    rows.append(row_dict)
            wb.close()
            logger.info(f"Read {len(rows)} CWC rows from Excel: {path}")
            return rows

        return []

    def _find_column(self, columns: List[str], target: str) -> Optional[str]:
        """Find a column matching a target pattern (case-insensitive)."""
        target_lower = target.lower()
        for col in columns:
            if target_lower in col.lower():
                return col
        return None

    def parse(self, raw_data: List[Dict]) -> List[CanonicalRecord]:
        """Parse CWC data rows into canonical records."""
        records = []
        if not raw_data:
            return records

        columns = list(raw_data[0].keys()) if raw_data else []

        # Find metadata columns
        station_col = self._find_column(columns, "station") or self._find_column(columns, "site")
        date_col = self._find_column(columns, "date")
        lat_col = self._find_column(columns, "lat")
        lon_col = self._find_column(columns, "lon") or self._find_column(columns, "long")
        basin_col = self._find_column(columns, "basin")
        river_col = self._find_column(columns, "river")
        state_col = self._find_column(columns, "state")

        # Detect parameter columns
        param_columns = {}
        for col in columns:
            col_lower = col.lower().strip()
            for pattern, (canon_param, canon_unit) in _CWC_COLUMN_MAP.items():
                if pattern in col_lower:
                    param_columns[col] = (canon_param, canon_unit)
                    break

        logger.info(f"CWC detected params: {list(param_columns.keys())}")

        for row in raw_data:
            station_name = self.clean_string(row.get(station_col, "")) if station_col else ""
            if not station_name:
                continue

            sample_date = self.safe_date(row.get(date_col, "")) if date_col else None
            if sample_date is None:
                continue

            lat = self.safe_float(row.get(lat_col, "")) if lat_col else None
            lon = self.safe_float(row.get(lon_col, "")) if lon_col else None
            river = self.clean_string(row.get(river_col, "")) if river_col else ""
            state = self.clean_string(row.get(state_col, "")) if state_col else ""

            # Generate a station ID from the name
            station_id = f"cwc_{station_name.lower().replace(' ', '_')[:30]}"

            for col, (canon_param, canon_unit) in param_columns.items():
                value = self.safe_float(row.get(col, ""))
                if value is None:
                    continue

                # Detect if discharge is in cusecs (common in CWC)
                actual_unit = canon_unit
                if canon_param == "discharge" and "cusec" in col.lower():
                    actual_unit = "ft³/s"

                records.append(CanonicalRecord(
                    station_id=station_id,
                    station_name=station_name,
                    river=river,
                    state=state,
                    latitude=lat or 0.0,
                    longitude=lon or 0.0,
                    sample_date=sample_date,
                    parameter=canon_param,
                    value=value,
                    unit=actual_unit,
                    source=self.source_id,
                    source_station_id=station_name,
                    source_parameter=col,
                    source_unit=actual_unit,
                    source_value=value,
                ))

        return records
